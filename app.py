from flask import Flask, render_template, request, redirect, url_for, session, flash, send_from_directory, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import sqlite3, os, secrets, mimetypes, smtplib, hashlib, re
from email.message import EmailMessage
from datetime import datetime, timedelta
from functools import wraps

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
# On hosts with a persistent disk, set DAILYUPDATES_DATA_DIR to its mount path
# (Render example: /var/data). Locally, data stays inside the project folder.
DATA_DIR = os.path.abspath(os.environ.get("DAILYUPDATES_DATA_DIR", BASE_DIR))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "dailyupdates.db")
UPLOAD_FOLDER = os.path.join(DATA_DIR, "uploads")
USER_FILES_ROOT = os.path.join(DATA_DIR, "user_files")
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "pdf", "txt", "doc", "docx", "xlsx"}
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(USER_FILES_ROOT, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("DAILYUPDATES_SECRET") or os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB per request

def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con

def init_db():
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            email TEXT,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'user',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS password_reset_otps (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL UNIQUE,
            otp_hash TEXT NOT NULL,
            expires_at TEXT NOT NULL,
            attempts INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS updates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            body TEXT NOT NULL,
            category TEXT NOT NULL,
            priority TEXT NOT NULL,
            audience TEXT NOT NULL DEFAULT 'All Users',
            attachment TEXT,
            created_by INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY(created_by) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS update_reads (
            user_id INTEGER NOT NULL,
            update_id INTEGER NOT NULL,
            read_at TEXT NOT NULL,
            PRIMARY KEY(user_id, update_id),
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE,
            FOREIGN KEY(update_id) REFERENCES updates(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS user_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            owner_id INTEGER NOT NULL,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL UNIQUE,
            file_category TEXT NOT NULL,
            is_shared INTEGER NOT NULL DEFAULT 0,
            uploaded_at TEXT NOT NULL,
            FOREIGN KEY(owner_id) REFERENCES users(id) ON DELETE CASCADE
        );
        """)
        # Safe migration: add email to existing databases without deleting users.
        columns = {
            row["name"] for row in con.execute("PRAGMA table_info(users)").fetchall()
        }
        if "email" not in columns:
            con.execute("ALTER TABLE users ADD COLUMN email TEXT")

        # Email addresses must be unique, ignoring case. Existing users can have NULL email.
        con.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS idx_users_email_unique
            ON users(email COLLATE NOCASE)
            WHERE email IS NOT NULL AND TRIM(email) != ''
        """)

        admin = con.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
        if not admin:
            con.execute(
                "INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)",
                ("admin", generate_password_hash("Admin@123"), "admin", datetime.now().isoformat(timespec="seconds"))
            )

def send_password_reset_otp(destination_email, otp):
    """Send a password reset OTP using SMTP environment settings."""
    smtp_host = os.environ.get("SMTP_HOST", "").strip()
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER", "").strip()
    smtp_password = os.environ.get("SMTP_PASSWORD", "")
    from_email = os.environ.get("SMTP_FROM", "").strip() or smtp_user

    if not all([smtp_host, smtp_user, smtp_password, from_email]):
        app.logger.error("Password reset email is not configured: SMTP settings are missing.")
        return False

    message = EmailMessage()
    message["Subject"] = "DailyUpdates password reset code"
    message["From"] = from_email
    message["To"] = destination_email
    message.set_content(
        f"Your DailyUpdates password reset OTP is: {otp}\n\n"
        "This code expires in 10 minutes. Do not share it with anyone. "
        "If you did not request a password reset, ignore this email."
    )

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=15) as server:
            server.starttls()
            server.login(smtp_user, smtp_password)
            server.send_message(message)
        return True
    except Exception:
        app.logger.exception("Could not send password reset email.")
        return False


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("Admin access is required.", "error")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)
    return wrapped

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route("/")
def home():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        if len(username) < 3 or len(username) > 80:
            flash("Username must be between 3 and 80 characters.", "error")
            return render_template("register.html")

        # Basic server-side email validation; the database enforces uniqueness.
        if (
            len(email) > 254
            or email.count("@") != 1
            or not email.split("@", 1)[0]
            or "." not in email.rsplit("@", 1)[-1]
            or any(char.isspace() for char in email)
        ):
            flash("Enter a valid email address.", "error")
            return render_template("register.html")

        if len(password) < 8:
            flash("Password must be at least 8 characters.", "error")
            return render_template("register.html")

        try:
            with db() as con:
                con.execute(
                    """INSERT INTO users
                       (username, email, password_hash, role, created_at)
                       VALUES (?, ?, ?, ?, ?)""",
                    (
                        username,
                        email,
                        generate_password_hash(password),
                        "user",
                        datetime.now().isoformat(timespec="seconds"),
                    ),
                )
            flash("Account created. You can now log in.", "success")
            return redirect(url_for("login"))
        except sqlite3.IntegrityError:
            flash("That username or email is already registered.", "error")

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        login_value = request.form.get("login", "").strip()
        password = request.form.get("password", "")

        with db() as con:
            user = con.execute(
                """SELECT * FROM users
                   WHERE username = ? COLLATE NOCASE
                      OR email = ? COLLATE NOCASE""",
                (login_value, login_value),
            ).fetchone()

        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect(url_for("dashboard"))

        flash("Invalid username/email or password.", "error")

    return render_template("login.html")


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        generic_message = (
            "If an account matches that email and email delivery is configured, "
            "a reset code will be sent."
        )

        with db() as con:
            user = con.execute(
                "SELECT id, email FROM users WHERE email = ? COLLATE NOCASE",
                (email,),
            ).fetchone()

            if user and user["email"]:
                otp = f"{secrets.randbelow(1_000_000):06d}"
                otp_hash = hashlib.sha256(otp.encode("utf-8")).hexdigest()
                now = datetime.now()
                expires = (now + timedelta(minutes=10)).isoformat(timespec="seconds")

                con.execute(
                    """INSERT INTO password_reset_otps
                       (user_id, otp_hash, expires_at, attempts, created_at)
                       VALUES (?, ?, ?, 0, ?)
                       ON CONFLICT(user_id) DO UPDATE SET
                           otp_hash=excluded.otp_hash,
                           expires_at=excluded.expires_at,
                           attempts=0,
                           created_at=excluded.created_at""",
                    (
                        user["id"],
                        otp_hash,
                        expires,
                        now.isoformat(timespec="seconds"),
                    ),
                )

                if not send_password_reset_otp(user["email"], otp):
                    app.logger.warning("Password reset email delivery failed.")

        flash(generic_message, "info")
        return redirect(url_for("reset_password", email=email))

    return render_template("forgot_password.html")


@app.route("/reset-password", methods=["GET", "POST"])
def reset_password():
    email = request.args.get("email", "").strip().lower()

    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        otp = request.form.get("otp", "").strip()
        new_password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if len(new_password) < 8:
            flash("Password must be at least 8 characters.", "error")
            return render_template("reset_password.html", email=email)

        if new_password != confirm_password:
            flash("Passwords do not match.", "error")
            return render_template("reset_password.html", email=email)

        if not re.fullmatch(r"\d{6}", otp):
            flash("Enter the 6-digit OTP sent to your email.", "error")
            return render_template("reset_password.html", email=email)

        with db() as con:
            user = con.execute(
                "SELECT id FROM users WHERE email = ? COLLATE NOCASE",
                (email,),
            ).fetchone()

            if not user:
                flash("Invalid or expired OTP. Request a new code.", "error")
                return render_template("reset_password.html", email=email)

            reset_record = con.execute(
                "SELECT * FROM password_reset_otps WHERE user_id = ?",
                (user["id"],),
            ).fetchone()

            if not reset_record:
                flash("Invalid or expired OTP. Request a new code.", "error")
                return render_template("reset_password.html", email=email)

            now = datetime.now()

            if reset_record["attempts"] >= 5:
                con.execute(
                    "DELETE FROM password_reset_otps WHERE user_id = ?",
                    (user["id"],),
                )
                flash("Too many attempts. Request a new OTP.", "error")
                return render_template("reset_password.html", email=email)

            if now > datetime.fromisoformat(reset_record["expires_at"]):
                con.execute(
                    "DELETE FROM password_reset_otps WHERE user_id = ?",
                    (user["id"],),
                )
                flash("OTP expired. Request a new code.", "error")
                return render_template("reset_password.html", email=email)

            provided_hash = hashlib.sha256(otp.encode("utf-8")).hexdigest()
            if not secrets.compare_digest(provided_hash, reset_record["otp_hash"]):
                con.execute(
                    "UPDATE password_reset_otps SET attempts = attempts + 1 WHERE user_id = ?",
                    (user["id"],),
                )
                flash("Invalid OTP.", "error")
                return render_template("reset_password.html", email=email)

            con.execute(
                "UPDATE users SET password_hash = ? WHERE id = ?",
                (generate_password_hash(new_password), user["id"]),
            )
            con.execute(
                "DELETE FROM password_reset_otps WHERE user_id = ?",
                (user["id"],),
            )

        flash("Password reset successfully. Please log in.", "success")
        return redirect(url_for("login"))

    return render_template("reset_password.html", email=email)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/dashboard")
@login_required
def dashboard():
    q = request.args.get("q", "").strip()
    category = request.args.get("category", "").strip()
    sql = """SELECT u.*, users.username AS author,
             CASE WHEN r.user_id IS NULL THEN 0 ELSE 1 END AS is_read
             FROM updates u JOIN users ON users.id=u.created_by
             LEFT JOIN update_reads r ON r.update_id=u.id AND r.user_id=?"""
    params = [session["user_id"]]
    if session.get("role") != "admin":
        sql += " WHERE (u.audience='All Users' OR u.audience=?)"
        params.append(session["username"])
    else:
        sql += " WHERE 1=1"
    if q:
        sql += " AND (u.title LIKE ? OR u.body LIKE ?)"
        params.extend([f"%{q}%", f"%{q}%"])
    if category:
        sql += " AND u.category=?"
        params.append(category)
    sql += " ORDER BY u.id DESC"
    with db() as con:
        updates = con.execute(sql, params).fetchall()
    categories = ["Announcement", "Meeting", "HR Notice", "System Update", "Training", "General", "Other"]
    return render_template("dashboard.html", updates=updates, categories=categories, q=q, category=category)

@app.route("/updates/create", methods=["GET", "POST"])
@admin_required
def create_update():
    categories = ["Announcement", "Meeting", "HR Notice", "System Update", "Training", "General", "Other"]
    priorities = ["Low", "Normal", "High", "Critical"]
    if request.method == "POST":
        title = request.form.get("title", "").strip()
        body = request.form.get("body", "").strip()
        category = request.form.get("category", "General")
        priority = request.form.get("priority", "Normal")
        audience = request.form.get("audience", "All Users")
        file = request.files.get("attachment")
        attachment = None
        if not title or not body:
            flash("Title and description are required.", "error")
            return render_template("create_update.html", categories=categories, priorities=priorities)
        if file and file.filename:
            if not allowed_file(file.filename):
                flash("File type not allowed. Use an image, PDF, TXT, DOC or DOCX.", "error")
                return render_template("create_update.html", categories=categories, priorities=priorities)
            safe_name = secure_filename(file.filename)
            unique_name = f"{secrets.token_hex(8)}_{safe_name}"
            file.save(os.path.join(UPLOAD_FOLDER, unique_name))
            attachment = unique_name
        with db() as con:
            con.execute("""INSERT INTO updates(title,body,category,priority,audience,attachment,created_by,created_at)
                           VALUES(?,?,?,?,?,?,?,?)""",
                        (title, body, category, priority, audience, attachment, session["user_id"],
                         datetime.now().isoformat(timespec="seconds")))
        flash("Update published and saved successfully.", "success")
        return redirect(url_for("dashboard"))
    return render_template("create_update.html", categories=categories, priorities=priorities)

@app.route("/updates/<int:update_id>")
@login_required
def view_update(update_id):
    with db() as con:
        item = con.execute("""SELECT u.*, users.username AS author FROM updates u
                              JOIN users ON users.id=u.created_by WHERE u.id=?""", (update_id,)).fetchone()
        if not item:
            flash("Update not found.", "error")
            return redirect(url_for("dashboard"))
        if session.get("role") != "admin" and item["audience"] not in ("All Users", session["username"]):
            flash("You do not have permission to view this update.", "error")
            return redirect(url_for("dashboard"))
        con.execute("INSERT OR REPLACE INTO update_reads(user_id,update_id,read_at) VALUES(?,?,?)",
                    (session["user_id"], update_id, datetime.now().isoformat(timespec="seconds")))
    return render_template("view_update.html", item=item)

@app.route("/uploads/<path:filename>")
@login_required
def uploaded_file(filename):
    # Only serve files that belong to an update visible to this user.
    with db() as con:
        row = con.execute("SELECT audience FROM updates WHERE attachment=?", (filename,)).fetchone()
    if not row:
        return "File not found", 404
    if session.get("role") != "admin" and row["audience"] not in ("All Users", session["username"]):
        return "Forbidden", 403
    return send_from_directory(UPLOAD_FOLDER, filename, as_attachment=False)

@app.route("/updates/<int:update_id>/delete", methods=["POST"])
@admin_required
def delete_update(update_id):
    with db() as con:
        row = con.execute("SELECT attachment FROM updates WHERE id=?", (update_id,)).fetchone()
        con.execute("DELETE FROM updates WHERE id=?", (update_id,))
    if row and row["attachment"]:
        try:
            os.remove(os.path.join(UPLOAD_FOLDER, row["attachment"]))
        except OSError:
            pass
    flash("Update deleted.", "success")
    return redirect(url_for("dashboard"))

@app.route("/users")
@admin_required
def users():
    with db() as con:
        all_users = con.execute("SELECT id,username,role,created_at FROM users ORDER BY id").fetchall()
    return render_template("users.html", users=all_users)

@app.route("/notes", methods=["GET", "POST"])
@login_required
def notes():
    if request.method == "POST":
        title = request.form.get("title", "").strip() or "Untitled note"
        body = request.form.get("body", "").strip()
        note_id = request.form.get("note_id", "").strip()
        if title == "Untitled note" and not body:
            flash("Write a title or note before saving.", "error")
            return redirect(url_for("notes"))
        now = datetime.now().isoformat(timespec="seconds")
        with db() as con:
            if note_id:
                con.execute("UPDATE notes SET title=?, body=?, updated_at=? WHERE id=? AND user_id=?",
                            (title, body, now, note_id, session["user_id"]))
                flash("Note updated and saved.", "success")
            else:
                con.execute("INSERT INTO notes(user_id,title,body,updated_at) VALUES(?,?,?,?)",
                            (session["user_id"], title, body, now))
                flash("Note saved successfully.", "success")
        return redirect(url_for("notes"))
    edit_id = request.args.get("edit", type=int)
    with db() as con:
        edit_note = con.execute("SELECT * FROM notes WHERE id=? AND user_id=?",
                                (edit_id, session["user_id"])).fetchone() if edit_id else None
        all_notes = con.execute("SELECT * FROM notes WHERE user_id=? ORDER BY updated_at DESC",
                                (session["user_id"],)).fetchall()
    return render_template("notes.html", notes=all_notes, edit_note=edit_note)

@app.route("/notes/<int:note_id>/delete", methods=["POST"])
@login_required
def delete_note(note_id):
    with db() as con:
        con.execute("DELETE FROM notes WHERE id=? AND user_id=?", (note_id, session["user_id"]))
    flash("Note deleted.", "success")
    return redirect(url_for("notes"))


@app.route("/files", methods=["GET", "POST"])
@login_required
def my_files():
    allowed = {"pdf", "doc", "docx", "txt", "xlsx", "png", "jpg", "jpeg", "gif", "webp"}
    if request.method == "POST":
        file = request.files.get("file")
        shared = 1 if session.get("role") == "admin" and request.form.get("is_shared") == "yes" else 0
        if not file or not file.filename:
            flash("Choose a file to upload.", "error")
            return redirect(url_for("my_files"))
        if "." not in file.filename or file.filename.rsplit(".", 1)[1].lower() not in allowed:
            flash("Unsupported file type. Use PDF, TXT, XLSX, DOC, DOCX, or an image.", "error")
            return redirect(url_for("my_files"))
        original_name = secure_filename(file.filename)
        ext = original_name.rsplit(".", 1)[1].lower()
        category = "images" if ext in {"png", "jpg", "jpeg", "gif", "webp"} else "documents"
        # Keep a separate folder for each owner and category.
        owner_folder = os.path.join(USER_FILES_ROOT, str(session["user_id"]), category)
        os.makedirs(owner_folder, exist_ok=True)
        stored_name = f"{secrets.token_hex(12)}_{original_name}"
        file.save(os.path.join(owner_folder, stored_name))
        with db() as con:
            con.execute("""INSERT INTO user_files(owner_id, original_name, stored_name, file_category, is_shared, uploaded_at)
                           VALUES(?,?,?,?,?,?)""",
                        (session["user_id"], original_name, stored_name, category, shared,
                         datetime.now().isoformat(timespec="seconds")))
        flash("File uploaded successfully.", "success")
        return redirect(url_for("my_files"))

    category_filter = request.args.get("category", "")
    with db() as con:
        if session.get("role") == "admin":
            # Admins can see all files uploaded by all users.
            rows = con.execute("""SELECT f.*, u.username AS owner_name FROM user_files f
                                  JOIN users u ON u.id=f.owner_id
                                  ORDER BY f.uploaded_at DESC""").fetchall()
        else:
            rows = con.execute("""SELECT f.*, u.username AS owner_name FROM user_files f
                                  JOIN users u ON u.id=f.owner_id
                                  WHERE f.owner_id=? OR f.is_shared=1
                                  ORDER BY f.uploaded_at DESC""", (session["user_id"],)).fetchall()
    if category_filter in ("images", "documents"):
        rows = [r for r in rows if r["file_category"] == category_filter]
    return render_template("files.html", files=rows, category=category_filter)

@app.route("/files/<int:file_id>/download")
@login_required
def download_user_file(file_id):
    with db() as con:
        if session.get("role") == "admin":
            item = con.execute("SELECT * FROM user_files WHERE id=?", (file_id,)).fetchone()
        else:
            item = con.execute(
                "SELECT * FROM user_files WHERE id=? AND (owner_id=? OR is_shared=1)",
                (file_id, session["user_id"])
            ).fetchone()
    if not item:
        return "File not found or access denied.", 404
    owner_folder = os.path.join(USER_FILES_ROOT, str(item["owner_id"]), item["file_category"])
    path = os.path.join(owner_folder, item["stored_name"])
    if not os.path.isfile(path):
        return "File is missing from storage.", 404
    return send_from_directory(owner_folder, item["stored_name"], as_attachment=True,
                               download_name=item["original_name"])

@app.route("/files/<int:file_id>/open")
@login_required
def open_user_file(file_id):
    """Open supported documents in the browser after checking file permissions."""
    with db() as con:
        if session.get("role") == "admin":
            item = con.execute("SELECT * FROM user_files WHERE id=?", (file_id,)).fetchone()
        else:
            item = con.execute(
                "SELECT * FROM user_files WHERE id=? AND (owner_id=? OR is_shared=1)",
                (file_id, session["user_id"])
            ).fetchone()
    if not item:
        return "File not found or access denied.", 404

    owner_folder = os.path.join(USER_FILES_ROOT, str(item["owner_id"]), item["file_category"])
    path = os.path.join(owner_folder, item["stored_name"])
    if not os.path.isfile(path):
        return "File is missing from storage.", 404

    ext = item["original_name"].rsplit(".", 1)[-1].lower() if "." in item["original_name"] else ""
    if ext in {"png", "jpg", "jpeg", "gif", "webp"}:
        mimetype = mimetypes.guess_type(item["original_name"])[0] or "application/octet-stream"
        return send_file(path, mimetype=mimetype, as_attachment=False,
                         download_name=item["original_name"])
    if ext == "pdf":
        response = send_file(path, mimetype="application/pdf", as_attachment=False,
                             download_name=item["original_name"])
        return response
    if ext == "txt":
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read(500_000)
        return render_template("text_preview.html", file=item, content=content)
    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for worksheet in workbook.worksheets:
                rows = []
                for row in worksheet.iter_rows(min_row=1, max_row=200, max_col=20, values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append(values)
                sheets.append({"name": worksheet.title, "rows": rows})
            workbook.close()
            return render_template("excel_preview.html", file=item, sheets=sheets)
        except Exception:
            flash("This Excel file could not be previewed. Try downloading it.", "error")
            return redirect(url_for("my_files"))
    flash("Preview is available for PDF, TXT, and XLSX files. Download this file to open it.", "error")
    return redirect(url_for("my_files"))


@app.route("/files/<int:file_id>/edit", methods=["GET", "POST"])
@login_required
def edit_user_file(file_id):
    """Edit owned TXT, XLSX, and DOCX files in the browser."""
    with db() as con:
        item = con.execute(
            "SELECT * FROM user_files WHERE id=? AND owner_id=?",
            (file_id, session["user_id"])
        ).fetchone()
    if not item:
        flash("Only the file owner can edit this file.", "error")
        return redirect(url_for("my_files"))

    ext = item["original_name"].rsplit(".", 1)[-1].lower() if "." in item["original_name"] else ""
    if ext not in {"txt", "xlsx", "docx"}:
        flash("Browser editing is available for TXT, XLSX, and DOCX files.", "error")
        return redirect(url_for("my_files"))
    folder = os.path.join(USER_FILES_ROOT, str(item["owner_id"]), item["file_category"])
    path = os.path.join(folder, item["stored_name"])
    if not os.path.isfile(path):
        flash("File is missing from storage.", "error")
        return redirect(url_for("my_files"))

    if ext == "txt":
        # Avoid a seemingly endless browser load when a large log/text file is opened.
        # Editing is intentionally limited to 500 KB; the original file is never
        # truncated just because it is too large for the in-browser editor.
        max_text_bytes = 500 * 1024
        try:
            if request.method == "POST":
                content = request.form.get("content", "")
                if len(content.encode("utf-8")) > max_text_bytes:
                    flash("This editor supports TXT files up to 500 KB. Your file was not changed.", "error")
                    return redirect(url_for("my_files"))
                with open(path, "w", encoding="utf-8", newline="") as handle:
                    handle.write(content)
                flash("Text file saved successfully.", "success")
                return redirect(url_for("edit_user_file", file_id=file_id))

            if os.path.getsize(path) > max_text_bytes:
                flash("This TXT file is larger than 500 KB, so it was not opened in the browser editor. Download it to edit locally; the original file is unchanged.", "error")
                return redirect(url_for("my_files"))
            with open(path, "r", encoding="utf-8-sig", errors="replace") as handle:
                content = handle.read(max_text_bytes + 1)
            return render_template("edit_text.html", file=item, content=content)
        except (OSError, UnicodeError) as exc:
            app.logger.exception("TXT editor failed for file id %s", file_id)
            flash("Could not open or save this TXT file. Check the VS Code terminal for the error; your original file has been preserved if opening failed.", "error")
            return redirect(url_for("my_files"))

    if ext == "docx":
        try:
            from docx import Document
            document = Document(path)
        except Exception:
            flash("This DOCX file could not be opened. Make sure it is a valid Word document.", "error")
            return redirect(url_for("my_files"))
        if request.method == "POST":
            paragraphs = request.form.getlist("paragraphs")
            if len(paragraphs) != len(document.paragraphs):
                flash("Document structure changed. Reload the editor and try again.", "error")
                return redirect(url_for("edit_user_file", file_id=file_id))
            for paragraph, value in zip(document.paragraphs, paragraphs):
                paragraph.text = value
            try:
                document.save(path)
            except Exception:
                flash("Could not save the Word document. Close it in other applications and try again.", "error")
                return redirect(url_for("edit_user_file", file_id=file_id))
            flash("Word document saved successfully.", "success")
            return redirect(url_for("edit_user_file", file_id=file_id))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        return render_template("edit_docx.html", file=item, paragraphs=paragraphs)

    # XLSX: edit one worksheet at a time while preserving other sheets and workbook content.
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path)
    except Exception:
        flash("This Excel file could not be opened. Make sure it is a valid .xlsx workbook.", "error")
        return redirect(url_for("my_files"))
    sheet_name = request.values.get("sheet", workbook.sheetnames[0])
    if sheet_name not in workbook.sheetnames:
        sheet_name = workbook.sheetnames[0]
    worksheet = workbook[sheet_name]
    max_row = min(max(worksheet.max_row, 1), 200)
    max_col = min(max(worksheet.max_column, 1), 20)
    if request.method == "POST":
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                key = f"cell_{row}_{col}"
                if key in request.form:
                    value = request.form.get(key, "")
                    cell = worksheet.cell(row=row, column=col)
                    # Keep a simple number as numeric data; otherwise save as text.
                    if value == "":
                        cell.value = None
                    else:
                        try:
                            if value.strip() and value.strip()[0] not in "=+@":
                                number = float(value) if "." in value else int(value)
                                cell.value = number
                            else:
                                cell.value = value
                        except (ValueError, TypeError):
                            cell.value = value
        try:
            workbook.save(path)
        except Exception:
            workbook.close()
            flash("Could not save the Excel file. Close it in Excel if it is currently open, then try again.", "error")
            return redirect(url_for("edit_user_file", file_id=file_id, sheet=sheet_name))
        workbook.close()
        flash(f"Worksheet '{sheet_name}' saved successfully.", "success")
        return redirect(url_for("edit_user_file", file_id=file_id, sheet=sheet_name))
    rows = []
    for row in range(1, max_row + 1):
        values = []
        for col in range(1, max_col + 1):
            value = worksheet.cell(row=row, column=col).value
            values.append("" if value is None else str(value))
        rows.append(values)
    sheet_names = workbook.sheetnames[:]
    workbook.close()
    return render_template("edit_excel.html", file=item, rows=rows, sheet_names=sheet_names,
                           selected_sheet=sheet_name, max_col=max_col, max_row=max_row)


@app.route("/files/<int:file_id>/share", methods=["POST"])
@admin_required
def toggle_file_share(file_id):
    with db() as con:
        row = con.execute("SELECT owner_id, is_shared FROM user_files WHERE id=?", (file_id,)).fetchone()
        if not row:
            flash("File not found.", "error")
        elif row["owner_id"] != session["user_id"]:
            flash("Admins can share their own uploaded files. User files remain private.", "error")
        else:
            new_value = 0 if row["is_shared"] else 1
            con.execute("UPDATE user_files SET is_shared=? WHERE id=?", (new_value, file_id))
            flash("File sharing updated.", "success")
    return redirect(url_for("my_files"))

@app.route("/files/<int:file_id>/delete", methods=["POST"])
@login_required
def delete_user_file(file_id):
    with db() as con:
        item = con.execute("SELECT * FROM user_files WHERE id=?", (file_id,)).fetchone()
        if not item or (item["owner_id"] != session["user_id"] and session.get("role") != "admin"):
            flash("File not found or access denied.", "error")
            return redirect(url_for("my_files"))
        # Admin may manage shared files, but cannot delete another user's private file.
        if item["owner_id"] != session["user_id"] and not item["is_shared"]:
            flash("You can only manage shared files belonging to other users.", "error")
            return redirect(url_for("my_files"))
        con.execute("DELETE FROM user_files WHERE id=?", (file_id,))
    folder = os.path.join(USER_FILES_ROOT, str(item["owner_id"]), item["file_category"])
    try:
        os.remove(os.path.join(folder, item["stored_name"]))
    except OSError:
        pass
    flash("File deleted.", "success")
    return redirect(url_for("my_files"))

@app.errorhandler(413)
def too_large(_error):
    flash("Upload is too large. Maximum request size is 10 MB.", "error")
    return redirect(request.referrer or url_for("dashboard"))

init_db()

if __name__ == "__main__":
    # Debug is disabled by default; enable only for local development if needed.
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=os.environ.get("FLASK_DEBUG") == "1")
